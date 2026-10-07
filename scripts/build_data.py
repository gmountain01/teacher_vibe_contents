"""구글 드라이브의 엑셀 파일을 내려받아 data.json으로 바꿉니다.

사용법: python scripts/build_data.py <출력 경로>
환경 변수 DRIVE_FILE_ID 로 파일 ID를 바꿀 수 있습니다.
"""
import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime, timedelta, timezone

import openpyxl

FILE_ID = os.environ.get("DRIVE_FILE_ID", "1003PRolBSSJFwivZT3vd_Sprv-vR2UPv")
SHEETS = ["수업자동화", "업무자동화"]


def download(file_id: str) -> bytes:
    urls = [
        f"https://drive.google.com/uc?export=download&id={file_id}",
        f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx",
    ]
    last = None
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=60) as res:
                data = res.read()
            if data[:2] == b"PK":  # xlsx(zip) 파일인지 확인
                return data
            last = f"{url} 에서 엑셀이 아닌 응답을 받음 (공유 설정 확인 필요)"
        except Exception as e:  # noqa: BLE001
            last = f"{url}: {e}"
    raise SystemExit(f"엑셀을 내려받지 못했습니다. {last}")


def clean(v) -> str:
    if v is None:
        return ""
    s = re.sub(r"[ \t]*\n[ \t]*", " ", str(v).strip())
    return re.sub(r" {2,}", " ", s)


def parse(wb) -> dict:
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            raise SystemExit(f"'{name}' 탭을 찾지 못했습니다. 탭 이름을 확인하세요.")
        ws = wb[name]
        ideas, ref1, ref2, mode = [], [], [], None
        for r in ws.iter_rows(values_only=True):
            r = list(r) + [None] * (10 - len(r))
            a = r[0]
            if isinstance(a, str):
                if a.startswith("①"):
                    mode = 1; continue
                if a.startswith("②"):
                    mode = 2; continue
                if a.startswith("③"):
                    mode = 3; continue
                if a.strip() in ("출처 도서", "번호"):
                    continue
            if mode == 1 and a:
                ref1.append({"book": str(a).strip(), "order": clean(r[3]), "toc": str(r[4] or "").strip(), "why": clean(r[6])})
            elif mode == 2 and a:
                ref2.append({"book": clean(a), "apps": str(r[4] or "").strip()})
            elif mode == 3 and isinstance(a, (int, float)) and clean(r[3]):
                ideas.append({
                    "n": int(a), "cat": clean(r[1]), "stage": clean(r[2]), "name": clean(r[3]),
                    "scene": clean(r[4]), "tool": clean(r[5]), "who": clean(r[6]),
                    "link": clean(r[7]), "note": clean(r[8]),
                })
        out[name] = {"ideas": ideas, "ref1": ref1, "ref2": ref2}
    return out


def main():
    dest = sys.argv[1] if len(sys.argv) > 1 else "data.json"
    wb = openpyxl.load_workbook(io.BytesIO(download(FILE_ID)), data_only=True)
    sheets = parse(wb)
    kst = datetime.now(timezone(timedelta(hours=9)))
    payload = {"updated": kst.strftime("%Y-%m-%d %H:%M"), "sheets": sheets}
    os.makedirs(os.path.dirname(os.path.abspath(dest)), exist_ok=True)
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
    print({k: len(v["ideas"]) for k, v in sheets.items()}, "→", dest)


if __name__ == "__main__":
    main()
